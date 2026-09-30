from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


OUTPUT = Path(__file__).with_name("Agentic-Development-Team-Demo.pptx")

NAVY = RGBColor.from_string("164A61")
BLUE = RGBColor.from_string("04668C")
SKY = RGBColor.from_string("8DC4E6")
PAPER = RGBColor.from_string("F9F6EF")
SAND = RGBColor.from_string("D4CDBF")


def add_text(slide, text, x, y, width, height, size=22, color=NAVY, bold=False):
    box = slide.shapes.add_textbox(
        Inches(x), Inches(y), Inches(width), Inches(height)
    )
    box.text_frame.word_wrap = True
    paragraph = box.text_frame.paragraphs[0]
    paragraph.text = text
    paragraph.font.name = "Aptos"
    paragraph.font.size = Pt(size)
    paragraph.font.color.rgb = color
    paragraph.font.bold = bold
    return box


def new_slide(prs, title, subtitle):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = PAPER
    add_text(slide, title, 0.7, 0.45, 12, 0.7, size=30, bold=True)
    add_text(slide, subtitle, 0.72, 1.2, 11.8, 0.6, size=16, color=BLUE)
    line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(7.25), Inches(13.333), Inches(0.25)
    )
    line.fill.solid()
    line.fill.fore_color.rgb = BLUE
    line.line.fill.background()
    add_text(
        slide, "Development demo | Production remains manual",
        0.72, 6.8, 10.8, 0.3, size=12,
    )
    page = add_text(slide, str(len(prs.slides)), 12.0, 6.8, 0.55, 0.3, size=12)
    page.text_frame.paragraphs[0].alignment = PP_ALIGN.RIGHT
    return slide


def add_bullets(slide, items, top=2.0, height=2.8):
    box = slide.shapes.add_textbox(Inches(0.85), Inches(top), Inches(11.6), Inches(height))
    frame = box.text_frame
    frame.word_wrap = True
    for index, item in enumerate(items):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = f"•  {item}"
        paragraph.font.name = "Aptos"
        paragraph.font.size = Pt(20)
        paragraph.font.color.rgb = NAVY
        paragraph.space_after = Pt(12)


def add_card(slide, title, body, y=4.9):
    card = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.85), Inches(y), Inches(11.6), Inches(1.45)
    )
    card.fill.solid()
    card.fill.fore_color.rgb = SKY
    card.line.color.rgb = SAND
    add_text(slide, title, 1.1, y + 0.12, 11.1, 0.35, size=17, bold=True)
    add_text(slide, body, 1.1, y + 0.55, 11.1, 0.75, size=18)


def add_flow(slide, labels):
    gap = 0.2
    width = (11.6 - gap * (len(labels) - 1)) / len(labels)
    for index, label in enumerate(labels):
        x = 0.85 + index * (width + gap)
        shape = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(3.1),
            Inches(width), Inches(1.0),
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = NAVY if index % 2 == 0 else BLUE
        shape.line.fill.background()
        paragraph = shape.text_frame.paragraphs[0]
        paragraph.text = label
        paragraph.alignment = PP_ALIGN.CENTER
        paragraph.font.name = "Aptos"
        paragraph.font.size = Pt(18)
        paragraph.font.bold = True
        paragraph.font.color.rgb = PAPER


def build_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    prs.core_properties.title = "Hotel demo: ask, track, see the result"
    prs.core_properties.subject = "Simple product-owner walkthrough; see TEAM-DEMO-GUIDE.md"

    slide = new_slide(
        prs, "Hotel demo: ask, track, see the result",
        "5-10 minutes | Use an existing change rather than waiting for a new build",
    )
    add_text(
        slide, "Describe a guest need. Follow the work. See the development result.",
        0.85, 2.0, 11.6, 0.8, size=24,
    )
    add_flow(slide, ["Ask", "Boards", "Cloud agent", "Checks", "Development"])
    add_card(
        slide, "One simple story",
        "Azure Boards shows the requirement. GitHub shows the work. The Hotel shows the result.",
    )

    slide = new_slide(prs, "1. Show the Hotel", "The guest experience comes first")
    add_bullets(
        slide,
        [
            "Open the development site and choose a hotel.",
            "Select future dates and guests, check availability, and show the total.",
            "Stop before Confirm booking: shared data does not need another reservation.",
        ],
    )
    add_card(slide, "If the site is unavailable", "Show existing cloud screenshots as recorded evidence.")

    slide = new_slide(prs, "2. Ask a knowledge question", "Show a cited answer, not just a confident answer")
    add_bullets(
        slide,
        [
            "Ask for the booking rules in plain language, with sources and revision.",
            "Repository knowledge works now; Search MCP needs a verified connection.",
            "For MCP, show the actual tool call and passage. Otherwise say it used repository files.",
        ],
    )
    add_card(
        slide, "Copy this prompt",
        "Read the repository knowledge. Explain our booking rules and cite the documents and revision. Do not change anything.",
    )

    slide = new_slide(prs, "3. Agree one small improvement", "No technical commands for the product owner")
    add_bullets(
        slide,
        [
            "Describe one new improvement and agree how to know it works.",
            "Say: That looks good. Create the story.",
            "Default: sample-project User Story, New, no github-synced tag.",
            "Scheduled intake assigns a cloud developer; pickup can be delayed.",
        ],
    )
    add_card(slide, "Read-only replay?", "Skip story creation. It starts real work.")

    slide = new_slide(prs, "4. Follow one existing change", "Story -> linked issue -> pull request")
    add_bullets(
        slide,
        [
            "Validation checks build and tests; Copilot reviews the exact revision.",
            "QA saves workflow evidence; the custom QA agent is not yet automated.",
            "One review alternative is normally skipped. It is not duplicate work.",
            "If checks are still running or fail, do not call the change delivered.",
        ],
    )
    add_card(slide, "Where to look", "Use Boards and GitHub. Agent Operations is not a mirror of every delivery step.")

    slide = new_slide(prs, "5. Show the development result", "No extra human release approval for this demo")
    add_bullets(
        slide,
        [
            "Once required checks pass, the coordinator can merge without asking again.",
            "Merge starts development deployment. Show the result only after success.",
            "Workflow execution approval is separate. Production stays manual.",
            "Prove MCP chat retrieval first; verify development-agent access separately.",
        ],
    )
    add_card(
        slide, "Before presenting",
        "Read TEAM-DEMO-GUIDE.md for live links and readiness notes. Pending PRs are not deployed features.",
    )

    prs.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build_deck())
